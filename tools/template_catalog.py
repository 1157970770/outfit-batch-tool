#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""本机视频标签项目与 XLSX 交接清单。"""
from __future__ import annotations

import hashlib
import json
import os
import shutil
import subprocess
import uuid
from pathlib import Path
from typing import Any

VIDEO_EXT = {".mp4", ".mov", ".m4v"}
FINGERPRINT_CHUNK = 1 << 20


def _probe_video(path: Path) -> dict:
    """本地读取时长；不导入生产应用模块。"""
    exe = shutil.which("ffprobe")
    if not exe:
        fallback = Path(r"D:\ffmpeg\ffmpeg-9.0.1-essentials_build\bin\ffprobe.exe")
        exe = str(fallback) if fallback.is_file() else ""
    if not exe:
        return {"warn": "未找到 ffprobe；标签仍可保存，时长留空"}
    try:
        result = subprocess.run(
            [exe, "-v", "error", "-show_entries", "format=duration",
             "-of", "default=noprint_wrappers=1:nokey=1", str(path)],
            capture_output=True, text=True, timeout=30,
        )
        duration = float(result.stdout.strip()) if result.returncode == 0 else None
    except (OSError, subprocess.SubprocessError, ValueError):
        duration = None
    return {"duration": round(duration, 2)} if duration and duration > 0 else {"warn": "无法读取视频时长"}


def _file_fingerprint(path: Path) -> str:
    """轻量指纹用于本机文件变更检测，不作为安全哈希。"""
    size = path.stat().st_size
    digest = hashlib.sha256()
    digest.update(str(size).encode("ascii"))
    with path.open("rb") as stream:
        digest.update(stream.read(FINGERPRINT_CHUNK))
        if size > 2 * FINGERPRINT_CHUNK:
            stream.seek(-FINGERPRINT_CHUNK, 2)
            digest.update(stream.read(FINGERPRINT_CHUNK))
    return digest.hexdigest()[:16]

DEFAULT_TAGS = [
    "适合秋冬", "适合春夏", "近景", "中远景",
    "手串展示", "A场景", "B场景", "C场景",
]
VIDEO_SHEET = "视频标签清单"
TAGS_SHEET = "标签字典"
PROJECT_SHEET = "项目"


class CatalogError(ValueError):
    pass


def library_path() -> Path:
    """全局标签库：跨文件夹、跨次启动保留。"""
    override = (os.getenv("OUTFIT_TAG_LIBRARY") or "").strip()
    if override:
        return Path(override)
    base = (os.getenv("APPDATA") or os.getenv("LOCALAPPDATA") or "").strip()
    if base:
        return Path(base) / "OutfitVideoTagger" / "tags.json"
    return Path.home() / ".outfit-video-tagger" / "tags.json"


def load_tag_library(path: str | Path | None = None) -> list[str]:
    """读取标签库；只有文件不存在/损坏时才用默认标签，已删除的默认标签不会复活。"""
    target = Path(path) if path else library_path()
    if not target.is_file():
        return list(DEFAULT_TAGS)
    try:
        data = json.loads(target.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError):
        return list(DEFAULT_TAGS)
    values = data.get("tags") if isinstance(data, dict) else data
    if not isinstance(values, list):
        return list(DEFAULT_TAGS)
    tags: list[str] = []
    for value in values:
        try:
            name = validate_tag_name(value)
        except CatalogError:
            continue
        if name.casefold() not in {tag.casefold() for tag in tags}:
            tags.append(name)
    return tags


def state_path() -> Path:
    """工作状态：上次项目、最近文件夹，供退出后继续。"""
    return library_path().with_name("state.json")


def load_state(path: str | Path | None = None) -> dict:
    target = Path(path) if path else state_path()
    if not target.is_file():
        return {}
    try:
        data = json.loads(target.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError):
        return {}
    return data if isinstance(data, dict) else {}


def save_state(state: dict, path: str | Path | None = None) -> Path:
    target = Path(path) if path else state_path()
    target.parent.mkdir(parents=True, exist_ok=True)
    tmp = target.with_suffix(target.suffix + ".tmp")
    tmp.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(tmp, target)
    return target


def save_tag_library(tags: list[str], path: str | Path | None = None) -> Path:
    target = Path(path) if path else library_path()
    target.parent.mkdir(parents=True, exist_ok=True)
    payload = {"version": 1, "tags": list(tags)}
    tmp = target.with_suffix(target.suffix + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(tmp, target)
    return target


def add_library_tag(tags: list[str], value: str) -> str:
    name = validate_tag_name(value)
    if name.casefold() in {tag.casefold() for tag in tags}:
        raise CatalogError("这个标签已经存在")
    tags.append(name)
    return name


def new_project(root: str | Path, progress=None) -> dict:
    folder = Path(root).expanduser().resolve()
    if not folder.is_dir():
        raise CatalogError(f"视频文件夹不存在：{folder}")
    project = {
        "project_id": uuid.uuid4().hex[:12],
        "source_root": str(folder),
        "next_id": 1,
        "tags": [],
        "videos": [],
    }
    return scan_folder(folder, project, progress=progress)


def _next_id(project: dict) -> int:
    numeric = [int(str(v.get("id"))) for v in project.get("videos", [])
               if str(v.get("id", "")).isdigit()]
    return max(int(project.get("next_id", 1)), max(numeric, default=0) + 1)


def scan_folder(root: str | Path, project: dict | None = None, progress=None) -> dict:
    folder = Path(root).expanduser().resolve()
    if not folder.is_dir():
        raise CatalogError(f"视频文件夹不存在：{folder}")
    result = dict(project or {})
    result.setdefault("project_id", uuid.uuid4().hex[:12])
    result.setdefault("tags", [])
    result["source_root"] = str(folder)
    previous = {str(row.get("relative_path")): dict(row) for row in result.get("videos", [])}
    found: set[str] = set()
    rows: list[dict] = []
    next_id = _next_id(result)
    paths = sorted((p for p in folder.rglob("*") if p.is_file() and p.suffix.lower() in VIDEO_EXT),
                   key=lambda p: p.relative_to(folder).as_posix().casefold())
    total_paths = len(paths)
    for index, path in enumerate(paths, 1):
        if progress:
            progress(index - 1, total_paths, path.name)
        relative = path.relative_to(folder).as_posix()
        found.add(relative)
        old = previous.get(relative)
        info = _probe_video(path)
        fingerprint = _file_fingerprint(path)
        if old:
            row = old
            old_fingerprint = str(row.get("fingerprint") or "")
            changed_now = bool(old_fingerprint and old_fingerprint != fingerprint)
            row["changed"] = bool(row.get("changed")) or changed_now
            if changed_now and not row.get("previous_fingerprint"):
                row["previous_fingerprint"] = old_fingerprint
            row["id"] = str(row.get("id") or f"{next_id:03d}")
        else:
            row = {"id": f"{next_id:03d}", "tags": [], "changed": False}
            next_id += 1
        row.update({
            "relative_path": relative,
            "name": path.name,
            "size_bytes": path.stat().st_size,
            "size_mb": round(path.stat().st_size / 1024 / 1024, 2),
            "duration": info.get("duration"),
            "fingerprint": fingerprint,
            "missing": False,
            "validation_warning": info.get("warn", ""),
        })
        row["tags"] = normalize_tags(row.get("tags", []), result["tags"])
        rows.append(row)
        if progress:
            progress(index, total_paths, path.name)
    # 保留 XLSX 中但当前文件夹暂缺的行，避免误删人工标签和 ID。
    for relative, old in previous.items():
        if relative not in found:
            old["missing"] = True
            old["changed"] = False
            rows.append(old)
    rows.sort(key=lambda row: (not row.get("missing", False), int(row.get("id", "0")) if str(row.get("id", "")).isdigit() else 0))
    result["videos"] = rows
    result["next_id"] = max(next_id, _next_id(result))
    return result


def validate_tag_name(value: str) -> str:
    name = str(value or "").strip()
    if not name:
        raise CatalogError("标签名称不能为空")
    if len(name) > 64:
        raise CatalogError("标签名称最多 64 个字符")
    if any(ch in name for ch in ("|", ";", "\r", "\n")):
        raise CatalogError("标签不能包含分隔符 |、; 或换行")
    return name


def normalize_tags(values: Any, known: list[str]) -> list[str]:
    if isinstance(values, str):
        values = values.split(";") if values else []
    if not isinstance(values, (list, tuple, set)):
        return []
    allowed = set(known)
    out: list[str] = []
    for value in values:
        label = str(value).strip()
        if label in allowed and label not in out:
            out.append(label)
    return out


def add_tag(project: dict, value: str) -> str:
    name = validate_tag_name(value)
    tags = list(project.get("tags") or [])
    if name.casefold() in {tag.casefold() for tag in tags}:
        raise CatalogError("这个标签已经存在")
    tags.append(name)
    project["tags"] = tags
    return name


def set_video_tags(project: dict, relative_path: str, selected: list[str]) -> None:
    tags = list(project.get("tags") or [])
    for value in selected:
        if value not in tags:
            raise CatalogError(f"标签不存在：{value}")
    for row in project.get("videos", []):
        if row.get("relative_path") == relative_path:
            row["tags"] = normalize_tags(selected, tags)
            return
    raise CatalogError("找不到指定视频")


def validate_project(project: dict, *, require_files: bool = True) -> list[str]:
    errors: list[str] = []
    root = Path(project.get("source_root") or "")
    ids: set[str] = set()
    paths: set[str] = set()
    known_tags = set(project.get("tags") or [])
    for row in project.get("videos", []):
        tid = str(row.get("id") or "")
        rel = str(row.get("relative_path") or "")
        if not tid.isdigit():
            errors.append(f"模板编号必须为数字：{tid or rel}")
        if not tid or tid in ids:
            errors.append(f"模板编号重复或为空：{tid or rel}")
        ids.add(tid)
        if not rel or rel in paths:
            errors.append(f"视频路径重复或为空：{rel}")
        paths.add(rel)
        if row.get("missing"):
            errors.append(f"本地视频不存在：{rel}")
            continue
        if row.get("changed"):
            errors.append(f"视频内容变化尚未确认：{rel}")
        source = root / Path(rel)
        if require_files and not source.is_file():
            errors.append(f"本地视频不存在：{rel}")
        elif require_files and row.get("fingerprint"):
            try:
                if _file_fingerprint(source) != row["fingerprint"]:
                    errors.append(f"视频内容在保存标签清单后发生变化：{rel}；请重新扫描并确认")
            except OSError:
                errors.append(f"无法读取视频：{rel}")
        row_tags = row.get("tags", [])
        if not row_tags:
            errors.append(f"尚未打标签：{rel}")
        unknown_tags = [tag for tag in row_tags if tag not in known_tags]
        if unknown_tags:
            errors.append(f"标签字典中不存在 {unknown_tags[0]}：{rel}")
    if not project.get("videos"):
        errors.append("项目中没有找到视频")
    return errors


def save_xlsx(project: dict, destination: str | Path) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise CatalogError("本机缺少 Excel 文件支持模块 openpyxl，请联系管理员部署") from None
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = VIDEO_SHEET
    headers = ["模板ID", "相对路径", "文件名", "时长(秒)", "大小(MB)", "文件指纹", "标签", "状态"]
    ws.append(headers)
    for row in project.get("videos", []):
        state = "文件不存在" if row.get("missing") else "文件已变化" if row.get("changed") else "待上传"
        ws.append([str(row.get("id", "")), str(row.get("relative_path", "")), str(row.get("name", "")),
                   row.get("duration"), row.get("size_mb"), row.get("fingerprint", ""),
                   ";".join(row.get("tags", [])), state])
    tags_ws = wb.create_sheet(TAGS_SHEET)
    tags_ws.append(["标签值"])
    for tag in project.get("tags", []):
        tags_ws.append([tag])
    meta = wb.create_sheet(PROJECT_SHEET)
    meta.append(["项目字段", "值"])
    meta.append(["项目ID", str(project.get("project_id") or "")])
    meta.append(["本机视频文件夹", str(project.get("source_root") or "")])
    meta.append(["下一个编号", int(project.get("next_id") or 1)])
    meta.sheet_state = "hidden"
    for sheet in (ws, tags_ws):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2768DF")
            cell.alignment = Alignment(vertical="center")
        for idx, column in enumerate(sheet.columns, 1):
            values = [len(str(c.value or "")) for c in list(column)[:300]]
            sheet.column_dimensions[get_column_letter(idx)].width = min(max(max(values, default=12) + 2, 12), 54)
    wb.save(path)
    return path


def load_xlsx(source: str | Path) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise CatalogError("本机缺少 Excel 文件支持模块 openpyxl") from None
    path = Path(source)
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise CatalogError(f"无法打开标签清单：{exc}") from None
    if VIDEO_SHEET not in wb.sheetnames or TAGS_SHEET not in wb.sheetnames:
        wb.close()
        raise CatalogError("Excel 缺少“视频标签清单”或“标签字典”工作表")
    tags_ws = wb[TAGS_SHEET]
    tags = [validate_tag_name(row[0]) for row in list(tags_ws.iter_rows(min_row=2, values_only=True)) if row and row[0]]
    if not tags:
        tags = list(DEFAULT_TAGS)
    meta_values = {}
    if PROJECT_SHEET in wb.sheetnames:
        meta = wb[PROJECT_SHEET]
        for key, value in meta.iter_rows(min_row=2, values_only=True):
            if key:
                meta_values[str(key)] = value
    ws = wb[VIDEO_SHEET]
    header = [str(c or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    column = {name: idx for idx, name in enumerate(header)}
    required = ("模板ID", "相对路径", "文件名", "标签")
    if any(name not in column for name in required):
        wb.close()
        raise CatalogError("视频标签清单缺少必要列（模板ID、相对路径、文件名、标签）")
    videos = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        tid = values[column["模板ID"]]
        relative = values[column["相对路径"]]
        if not tid or not relative:
            continue
        def get(name, default=None):
            idx = column.get(name)
            return values[idx] if idx is not None and idx < len(values) and values[idx] is not None else default
        rel = Path(str(relative).replace("\\", "/"))
        if rel.is_absolute() or ".." in rel.parts:
            wb.close()
            raise CatalogError(f"清单包含越界相对路径：{relative}")
        raw_tags = [part.strip() for part in str(get("标签", "")).split(";") if part.strip()]
        unknown = [tag for tag in raw_tags if tag not in tags]
        if unknown:
            wb.close()
            raise CatalogError(f"标签字典中不存在“{unknown[0]}”（视频 {rel.as_posix()}）")
        videos.append({"id": str(tid).strip(), "relative_path": rel.as_posix(),
                       "name": str(get("文件名", rel.name)), "duration": get("时长(秒)"),
                       "size_mb": get("大小(MB)", 0), "fingerprint": str(get("文件指纹", "")),
                       "tags": normalize_tags(raw_tags, tags), "missing": get("状态") == "文件不存在",
                       "changed": get("状态") == "文件已变化"})
    wb.close()
    source_root = str(meta_values.get("本机视频文件夹") or path.parent)
    project_id = str(meta_values.get("项目ID") or uuid.uuid4().hex[:12])
    try:
        next_id = int(meta_values.get("下一个编号") or (max((int(v["id"]) for v in videos if v["id"].isdigit()), default=0) + 1))
    except (TypeError, ValueError):
        next_id = 1
    return {"project_id": project_id, "source_root": source_root, "next_id": next_id,
            "tags": tags, "videos": videos, "xlsx_path": str(path)}
